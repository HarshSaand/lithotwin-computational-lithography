#!/usr/bin/env python3
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'; FIG=OUT/'figures'; m=json.loads((OUT/'metrics.json').read_text()); t=m['test']; op=m['ood_process']; og=m['ood_geometry']
NAVY='082B4C'; BLUE='2E74B5'; TEAL='008C95'; GOLD='D6A33A'; PALE='F2F4F7'; GRAY='536777'; WHITE='FFFFFF'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)
for name,size,color,bold,bef,aft,line in [('Normal',11,'263746',False,0,6,1.10),('Title',28,NAVY,True,0,8,1.0),('Subtitle',14,GRAY,False,0,12,1.0),('Heading 1',16,BLUE,True,16,8,1.0),('Heading 2',13,BLUE,True,12,6,1.0),('Heading 3',12,'1F4D78',True,8,4,1.0)]:
 s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=bold; s.paragraph_format.space_before=Pt(bef); s.paragraph_format.space_after=Pt(aft); s.paragraph_format.line_spacing=line
h=sec.header.paragraphs[0]; h.text='LITHOTWIN  |  TECHNICAL REPORT'; h.runs[0].font.size=Pt(8); h.runs[0].font.bold=True; h.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; f.add_run('Harsh Saand  |  Physics-generated prototype evidence  |  '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
for r in f.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

def shade(c,fill):
 pr=c._tc.get_or_add_tcPr(); x=pr.find(qn('w:shd')) or OxmlElement('w:shd'); x.set(qn('w:fill'),fill); pr.append(x) if x.getparent() is None else None
def set_table_geometry(tbl,widths):
 dxa=[int(w*1440) for w in widths]; pr=tbl._tbl.tblPr; tw=pr.find(qn('w:tblW')) or OxmlElement('w:tblW'); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),'9360'); pr.append(tw) if tw.getparent() is None else None
 ind=pr.find(qn('w:tblInd')) or OxmlElement('w:tblInd'); ind.set(qn('w:type'),'dxa'); ind.set(qn('w:w'),'120'); pr.append(ind) if ind.getparent() is None else None
 grid=tbl._tbl.tblGrid
 for x in list(grid): grid.remove(x)
 for w in dxa: c=OxmlElement('w:gridCol'); c.set(qn('w:w'),str(w)); grid.append(c)
 for row in tbl.rows:
  for cell,w in zip(row.cells,dxa):
   tcpr=cell._tc.get_or_add_tcPr(); tcw=tcpr.find(qn('w:tcW')) or OxmlElement('w:tcW'); tcw.set(qn('w:type'),'dxa'); tcw.set(qn('w:w'),str(w)); tcpr.append(tcw) if tcw.getparent() is None else None
   mar=tcpr.find(qn('w:tcMar')) or OxmlElement('w:tcMar')
   for tag,val in [('top',80),('bottom',80),('start',120),('end',120)]:
    e=mar.find(qn('w:'+tag)) or OxmlElement('w:'+tag); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e) if e.getparent() is None else None
   tcpr.append(mar) if mar.getparent() is None else None
def table(headers,rows,widths):
 tbl=doc.add_table(rows=1,cols=len(headers)); tbl.style='Table Grid'; tbl.alignment=WD_TABLE_ALIGNMENT.LEFT; tbl.autofit=False
 for i,x in enumerate(headers):
  c=tbl.rows[0].cells[i]; c.text=str(x); shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for z in c.paragraphs[0].runs: z.font.bold=True; z.font.color.rgb=RGBColor.from_string(WHITE); z.font.size=Pt(8.5)
 for j,row in enumerate(rows):
  cs=tbl.add_row().cells
  for i,x in enumerate(row):
   cs[i].text=str(x); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if j%2: shade(cs[i],PALE)
   for z in cs[i].paragraphs[0].runs: z.font.size=Pt(8.5)
 set_table_geometry(tbl,widths); p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
def pic(name,width,caption):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FIG/name),width=Inches(width)); q=doc.add_paragraph(caption); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_before=Pt(4); q.paragraph_format.space_after=Pt(4); q.runs[0].italic=True; q.runs[0].font.size=Pt(8.5); q.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
def callout(label,text):
 tbl=doc.add_table(rows=1,cols=1); c=tbl.cell(0,0); shade(c,'EAF4F4'); p=c.paragraphs[0]; a=p.add_run(label+'  '); a.bold=True; a.font.color.rgb=RGBColor.from_string(TEAL); p.add_run(text); set_table_geometry(tbl,[6.5]); doc.add_paragraph().paragraph_format.space_after=Pt(2)
def para(text): doc.add_paragraph(text)

# editorial_cover pattern with restrained engineering identity.
doc.add_paragraph('A*STAR IME PORTFOLIO PROTOTYPE').runs[0].font.color.rgb=RGBColor.from_string(TEAL)
doc.add_paragraph('LithoTwin',style='Title'); doc.add_paragraph('Conditional neural surrogate for computational lithography and resist-contour prediction',style='Subtitle'); doc.add_paragraph('HARSH SAAND',style='Heading 2'); doc.add_paragraph('Technical report  |  6 August 2026  |  Locally measured on Apple M3 Pro CPU')
pic('system_flow.png',4.3,'Implemented mask-to-contour workflow.')
callout('Evidence boundary','All quantitative results measure fidelity to the declared nine-source scalar optics and threshold-resist simulator. This is not production lithography, OPC, resist chemistry, or fab validation.')
doc.add_page_break(); doc.add_heading('Executive summary',1)
para(f"LithoTwin implements a complete AI-for-process-modeling workflow: procedural semiconductor mask geometries, nine-source scalar Abbe imaging, dose/focus/NA/threshold variation, geometry-grouped partitioning, a Gaussian imaging baseline, a conditional U-Net contour surrogate, engineering metrics, uncertainty stress tests, and an interactive process-window workbench. On 252 held-out cases from unseen base geometries, the U-Net achieved mean IoU {t['ai']['iou']:.3f}, Dice {t['ai']['dice']:.3f}, and symmetric edge distance {t['ai']['edge_distance_px']:.2f} pixels.")
callout('Measured contribution',f"The neural model improved held-out IoU from {t['baseline']['iou']:.3f} to {t['ai']['iou']:.3f} and Dice from {t['baseline']['dice']:.3f} to {t['ai']['dice']:.3f}. Separate process-range and unseen-contact-family OOD tests failed sharply, establishing a clear need to defer outside the validated domain.")
doc.add_heading('Computational lithography problem',1)
para('Lithography transfers a mask layout through an optical system and resist process into a printed contour. Small changes in dose, focus, numerical aperture, pitch, and threshold can change critical dimensions, merge neighbouring features, shorten line ends, or eliminate contacts. A conditional surrogate must therefore predict spatial contours while preserving process sensitivity and declaring where its training domain ends.')
pic('system_flow.png',4.6,'Figure 1. Implemented LithoTwin system flow.')

doc.add_page_break(); doc.add_heading('Reduced-order lithography physics',1)
doc.add_heading('Mask families',2); para('The deterministic generator creates 48 x 48 binary clips representing isolated lines, line ends, spaces, elbows, dense line arrays, two-line gaps, and circular contacts. Each base geometry receives multiple process conditions. Complete base IDs remain in a single partition, preventing near-identical recipe variants of one mask from leaking across train and test.')
doc.add_heading('Nine-source scalar Abbe approximation',2); para('For each source position, the mask spectrum is filtered by a shifted circular pupil. A quadratic phase term represents defocus. Source intensities are averaged, multiplied by relative exposure dose, and thresholded into a binary developed-resist region. The model is more process-aware than ordinary image blur, while remaining lightweight and auditable.')
table(['Stage','Inputs','Output','Declared simplification'],[('Mask rasterization','Family, width, pitch, location','Binary transmission mask','Procedural Manhattan clips'),('Optical imaging','Mask, NA, source positions, defocus, dose','Aerial intensity','Scalar model; nine source points'),('Resist development','Aerial image, threshold','Binary resist contour','No chemical diffusion or kinetics')],[1.3,1.8,1.5,1.9])
pic('prediction_panel.png',6.4,'Figure 2. Representative geometry-grouped test case from mask to signed contour error.')

doc.add_page_break(); doc.add_heading('Dataset and leakage-safe evaluation',1)
para('The seed-42 dataset contains 3,360 simulated mask/process cases. The primary process domain uses relative dose 0.78-1.22, defocus -1.15 to 1.15, NA 0.58-0.72, and threshold 0.25-0.48. Every tenth base geometry is held out for test and the next for validation. One extreme condition per non-contact base forms the process OOD set. The entire contact family is excluded from training and reserved as geometry OOD.')
table(['Partition','Cases','Purpose'],[('Train','2,016','Fit conditional contour model'),('Validation','252','Checkpoint selection'),('Held-out test','252','Unseen base geometries in primary process domain'),('Process OOD','360','Extreme dose and defocus'),('Geometry OOD','480','Entire unseen contact family')],[1.45,.85,4.2])
doc.add_heading('Metrics',2); table(['Metric','Engineering meaning'],[('IoU / Dice','Overlap of predicted and simulated developed resist'),('Symmetric edge distance','Average bidirectional contour displacement in pixels'),('Area error','Relative printed-area discrepancy; unstable for near-empty references'),('Brier score','Pixel-level probability error'),('Uncertainty-error correlation','Whether MC-dropout disagreement rises with contour error'),('Latency','Measured batched neural and reference-simulator CPU time')],[2.1,4.4])
para('Area percentage error becomes extremely large when the reference has only a few printed pixels. It is retained for transparency but is not treated as the primary metric; IoU, Dice, and edge distance are more stable for this dataset.')

doc.add_page_break(); doc.add_heading('AI methodology',1)
doc.add_heading('Gaussian baseline',2); para('The baseline applies a defocus-dependent Gaussian blur, multiplies by dose, and thresholds the result. It provides a meaningful low-cost approximation that lacks the shifted-pupil interference structure of the reference simulator.')
doc.add_heading('Conditional U-Net',2); para('The neural model receives five 48 x 48 channels: binary mask, relative dose, defocus, numerical aperture, and resist threshold. A compact encoder-decoder with skip connections predicts a continuous resist probability field. Training minimizes binary cross-entropy plus Dice loss with AdamW for 12 epochs; the lowest validation loss checkpoint is retained. Fixed seed 42 and CPU execution support reproducibility.')
table(['Component','Implementation','Purpose'],[('Input','Mask + four broadcast process channels','Condition spatial prediction on recipe'),('Encoder','Two convolution blocks with max pooling','Capture local and multi-scale geometry'),('Bottleneck','64-channel block with spatial dropout','Regularized latent representation'),('Decoder','Bilinear upsampling and skip connections','Recover contour localization'),('Output','One-channel logits','Resist probability at every pixel')],[1.2,2.5,2.8])
pic('process_window.png',6.3,'Figure 3. Physics and neural dose-focus process windows for a held-out geometry.')

doc.add_page_break(); doc.add_heading('Locally measured results',1)
table(['Metric','Gaussian baseline','Conditional U-Net'],[('Mean IoU',f"{t['baseline']['iou']:.3f}",f"{t['ai']['iou']:.3f}"),('Mean Dice',f"{t['baseline']['dice']:.3f}",f"{t['ai']['dice']:.3f}"),('Edge distance',f"{t['baseline']['edge_distance_px']:.2f} px",f"{t['ai']['edge_distance_px']:.2f} px"),('Pixel Brier score','Not probabilistic',f"{t['brier']:.4f}"),('Uncertainty-error correlation','Not available',f"{t['uncertainty_error_correlation']:.3f}")],[2.3,2.0,2.2])
pic('iou_comparison.png',5.8,'Figure 4. Mean contour IoU across interpolation and OOD regimes.')
pic('engineering_metrics.png',6.1,'Figure 5. Held-out contour errors and measured CPU latency.')
para(f"The U-Net substantially improves contour overlap in the validated regime. Batched neural inference measured {t['ai_us_per_clip']:.0f} us per clip, while the already-lightweight reference simulator measured {t['simulator_us_per_clip']:.0f} us per clip. The neural model is therefore not faster in this implementation, and no acceleration claim is made. The value demonstrated is conditional contour learning, process-window fidelity, and explicit domain testing.")

doc.add_page_break(); doc.add_heading('OOD and uncertainty findings',1)
table(['Evaluation regime','Cases','AI IoU','Baseline IoU','AI edge distance'],[('Held-out primary domain',t['n'],f"{t['ai']['iou']:.3f}",f"{t['baseline']['iou']:.3f}",f"{t['ai']['edge_distance_px']:.2f} px"),('Extreme process conditions',op['n'],f"{op['ai']['iou']:.3f}",f"{op['baseline']['iou']:.3f}",f"{op['ai']['edge_distance_px']:.2f} px"),('Unseen contact family',og['n'],f"{og['ai']['iou']:.3f}",f"{og['baseline']['iou']:.3f}",f"{og['ai']['edge_distance_px']:.2f} px")],[2.2,.7,1.0,1.2,1.4])
para(f"Process-range OOD IoU falls to {op['ai']['iou']:.3f}; unseen-contact IoU falls to {og['ai']['iou']:.3f}. The Gaussian baseline is stronger than the U-Net on unseen contacts, showing that generic optical smoothing can generalize better than learned morphology in this regime. MC-dropout uncertainty correlates with error at {og['uncertainty_error_correlation']:.3f} for unseen contacts but only {op['uncertainty_error_correlation']:.3f} for process OOD, so it is not a reliable universal detector.")
callout('Operational interpretation','Inputs outside the dose/focus range or from an unseen geometry family must be routed to the reference simulator. The Streamlit interface displays this warning directly.')
doc.add_heading('Engineer-facing interpretation',1)
para('Contour overlays localize edge movement and missing or extra resist. Dose-focus maps show whether the learned response follows the reference process window. Signed error panels distinguish over-print from under-print. The held-out and OOD tables quantify applicability instead of relying on visually attractive examples alone.')

doc.add_page_break(); doc.add_heading('Limitations and validity threats',1)
table(['Limitation','Impact on interpretation'],[('Physics-generated evidence','Validates surrogate methodology, not fab accuracy'),('Scalar nine-source optics','Omits vector polarization, calibrated aberrations, flare, and dense source integration'),('Threshold resist','Omits diffusion, chemical amplification, stochastic effects, and development kinetics'),('Procedural 48 x 48 masks','Underrepresent real GDS complexity and long-range context'),('Pixel grid','Limits subpixel critical-dimension and contour accuracy'),('MC-dropout uncertainty','Does not reliably detect every process-range shift'),('CPU implementation','Reference simulator is faster than this neural model'),('Single seed/checkpoint','No multi-seed confidence interval is claimed')],[2.1,4.4])
doc.add_heading('Conclusion',1)
para(f"LithoTwin demonstrates a complete and candid AI-for-computational-lithography prototype. On geometry-grouped held-out data, the conditional U-Net reaches IoU {t['ai']['iou']:.3f} and Dice {t['ai']['dice']:.3f}, materially outperforming the Gaussian baseline. Its process-window visualization and contour overlays are technically interpretable. The severe OOD degradation, incomplete uncertainty detection, and lack of CPU speedup are reported as findings rather than hidden.")
para('The resulting portfolio evidence spans optical simulation, structured semiconductor geometry generation, leakage-safe spatial learning, contour metrics, probability calibration, OOD analysis, interactive engineering visualization, testing, and reproducible reporting. Commercial or fab use would require calibrated optics, realistic resist physics, real layout clips, multi-seed validation, and measured wafer data; none is claimed here.')
doc.add_page_break(); doc.add_heading('References and model provenance',1)
table(['Reference','Use'],[('J. W. Goodman, Introduction to Fourier Optics, 3rd ed.','Scalar Fourier imaging foundations'),('A. K. Wong, Resolution Enhancement Techniques in Optical Lithography','Lithography process-window and imaging context'),('C. A. Mack, Fundamental Principles of Optical Lithography','Optical and resist modeling context'),('Repository source files','Exact simulator equations, seed, data splits, checkpoint, and metrics')],[4.4,2.1])
out=OUT/'lithotwin_report.docx'; doc.save(out); print(out)
